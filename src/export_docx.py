"""
Export `paper/PlateBack_Lite_Research_Paper.md` to a formatted Microsoft Word
document (`paper/PlateBack_Lite_Research_Paper.docx`), preserving headings,
inline formatting, tables, embedded figures, captions, and references.
"""

from __future__ import annotations

from pathlib import Path
import re
from typing import List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def _clean_latex_inline(text: str) -> str:
    """Convert common Markdown/LaTeX math tokens into clean Unicode for Word."""
    s = text
    replacements = {
        r"\text{MAE}": "MAE",
        r"\text{RMSE}": "RMSE",
        r"\text{SD}": "SD",
        r"\text{round}": "round",
        r"\text{km}^2": "km²",
        r"\text{M}": "M",
        r"\text{train}": "train",
        r"\text{test}": "test",
        r"\text{Served}": "Served",
        r"\text{Surplus}": "Surplus",
        r"\text{Total Simulated Surplus}": "Total Simulated Surplus",
        r"\text{Total Simulated Shortage}": "Total Simulated Shortage",
        r"\text{Demand Served (\%)}": "Demand Served (%)",
        r"\text{Relative Surplus Change vs. Baseline (\%)}": "Relative Surplus Change vs. Baseline (%)",
        r"\text{policy}": "policy",
        r"\text{baseline}": "baseline",
        r"\text{price\_diff}": "price_diff",
        r"\text{base\_price}": "base_price",
        r"\text{checkout\_price}": "checkout_price",
        r"\text{discount\_ratio}": "discount_ratio",
        r"\text{is\_discounted}": "is_discounted",
        r"\text{promo\_count}": "promo_count",
        r"\text{emailer\_for\_promotion}": "emailer_for_promotion",
        r"\text{homepage\_featured}": "homepage_featured",
        r"\text{week\_sin}": "week_sin",
        r"\text{week\_cos}": "week_cos",
        r"\text{op\_area}": "op_area",
        r"\text{week}": "week",
        r"\text{center\_id}": "center_id",
        r"\text{meal\_id}": "meal_id",
        r"\text{prior\_global\_cnt}": "prior_global_cnt",
        r"\text{hist\_meal\_mean}": "hist_meal_mean",
        r"\text{hist\_center\_mean}": "hist_center_mean",
        r"\text{num\_orders}": "num_orders",
        r"\mathcal{D}_{\text{train}}": "D_train",
        r"\mathcal{D}_{\text{test}}": "D_test",
        r"\mathcal{D}": "D",
        r"\mathbb{I}": "I",
        r"\downarrow": "↓",
        r"\uparrow": "↑",
        r"\le": "≤",
        r"\ge": "≥",
        r"\neq": "≠",
        r"\times": "×",
        r"\cdot": "·",
        r"\alpha": "α",
        r"\pi": "π",
        r"\sum": "∑",
        r"\sqrt": "√",
        r"\hat{D}": "D̂",
        r"\hat{y}": "ŷ",
        r"\bar{y}": "ȳ",
        r"\bar{D}": "D̄",
        r"R^2": "R²",
        r"\bmod": "mod",
        r"\qquad": "    ",
        r"\;": " ",
        r"\,": "",
        r"\!": "",
        r"\left(": "(",
        r"\right)": ")",
        r"\big(": "(",
        r"\big)": ")",
        r"\text{–}": "–",
        r"{,}": ",",
    }
    for k, v in replacements.items():
        s = s.replace(k, v)
    # Strip remaining $ delimiters
    s = s.replace("$$", "").replace("$", "")
    return s


def _add_formatted_runs(paragraph, text: str) -> None:
    """Parse **bold**, *italic*, and `code` spans and add styled runs to `paragraph`."""
    cleaned = _clean_latex_inline(text)
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    tokens = pattern.split(cleaned)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(tok)


def _render_markdown_table(doc: Document, table_lines: List[str]) -> None:
    """Convert Markdown pipe-table lines into a formatted Word table."""
    parsed_rows: List[List[str]] = []
    for line in table_lines:
        stripped = line.strip().strip("|")
        cells = [c.strip() for c in stripped.split("|")]
        # Skip separator row (e.g. :--- | :---:)
        if all(set(c).issubset({":", "-", " "}) for c in cells if c):
            continue
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    n_rows = len(parsed_rows)
    n_cols = max(len(r) for r in parsed_rows)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row_cells in enumerate(parsed_rows):
        for c_idx in range(n_cols):
            val = row_cells[c_idx] if c_idx < len(row_cells) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            _add_formatted_runs(p, val)
            for run in p.runs:
                run.font.size = Pt(9.5)
                if r_idx == 0:
                    run.bold = True

    doc.add_paragraph()


def export_markdown_to_docx(
    md_path: str | Path = "paper/PlateBack_Lite_Research_Paper.md",
    docx_path: str | Path = "paper/PlateBack_Lite_Research_Paper.docx",
    project_root: str | Path = ".",
) -> Path:
    """
    Convert the Markdown research manuscript into a formatted `.docx` file.
    """
    md_file = Path(md_path)
    out_file = Path(docx_path)
    root_dir = Path(project_root)

    lines = md_file.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # Configure base Normal style font
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    i = 0
    n = len(lines)
    in_math_block = False
    math_buffer: List[str] = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Skip horizontal rules
        if stripped == "---":
            i += 1
            continue

        # Handle multi-line $$ math blocks
        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 2:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(_clean_latex_inline(stripped))
            run.italic = True
            i += 1
            continue
        elif stripped == "$$":
            if not in_math_block:
                in_math_block = True
                math_buffer = []
            else:
                in_math_block = False
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(_clean_latex_inline(" ".join(math_buffer)))
                run.italic = True
            i += 1
            continue

        if in_math_block:
            math_buffer.append(stripped)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines: List[str] = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _render_markdown_table(doc, table_lines)
            continue

        # Handle Markdown Images ![alt](path)
        img_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if img_match:
            alt_text, rel_img_path = img_match.group(1), img_match.group(2)
            img_full = root_dir / rel_img_path
            if img_full.is_file():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(str(img_full), width=Inches(6.0))
            else:
                p_img = doc.add_paragraph()
                _add_formatted_runs(p_img, f"[{alt_text}: {rel_img_path}]")
            i += 1
            continue

        # Handle Headings
        if stripped.startswith("# "):
            heading_text = _clean_latex_inline(stripped[2:].strip())
            if heading_text == "1. Title":
                # Look ahead for actual paper title line
                i += 1
                while i < n and not lines[i].strip():
                    i += 1
                if i < n:
                    title_str = lines[i].strip().strip("*")
                    t_par = doc.add_heading(title_str, level=0)
                    t_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                i += 1
                continue
            doc.add_heading(heading_text, level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(_clean_latex_inline(stripped[4:].strip()), level=2)
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(_clean_latex_inline(stripped[3:].strip()), level=2)
            i += 1
            continue

        # Handle Bullet Lists
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_runs(p, stripped[2:].strip())
            i += 1
            continue

        # Handle Indented Sub-bullets
        if line.startswith("  - ") or line.startswith("   - "):
            p = doc.add_paragraph(style="List Bullet 2")
            _add_formatted_runs(p, stripped[2:].strip())
            i += 1
            continue

        # Handle Numbered Lists
        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            p = doc.add_paragraph()
            run_num = p.add_run(f"{num_match.group(1)}. ")
            run_num.bold = True
            _add_formatted_runs(p, num_match.group(2))
            i += 1
            continue

        # Handle Figure Captions (*Figure X ...*)
        if stripped.startswith("*Figure ") and stripped.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(_clean_latex_inline(stripped.strip("*")))
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        _add_formatted_runs(p, stripped)
        i += 1

    out_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_file))
    return out_file


if __name__ == "__main__":
    saved = export_markdown_to_docx()
    print(f"Exported Word document to: {saved.resolve()}")
